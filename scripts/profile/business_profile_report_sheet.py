"""Turn one or more business-profile eval reports into a spreadsheet.

``business_profile_eval run`` writes a report JSON per run under
logs/business-profile-eval/. Reading one in a terminal is fine for a single
number; comparing two runs, or working out which twelve cases to adjudicate,
is not. This builds an .xlsx with one tab per question a reader actually
asks, and the ``publish-eval-sheet`` skill (.claude/skills/) uploads it to
Google Drive as a native Sheet.

    python -m scripts.profile.business_profile_report_sheet \\
        logs/business-profile-eval/report-<a>.json [report-<b>.json ...] \\
        --out logs/business-profile-eval/eval-sheet.xlsx

Tabs:

- Summary      -- one column per run: headline search-addressable P/R/F1,
                  then every per-field metric the harness reports (accuracy
                  beside its majority baseline, coverage, accuracy when
                  committed on answerable cases, macro-F1, n). Two reports
                  side by side is the model comparison.
- Per-class    -- precision/recall/F1/support for every value of every
                  field, with `reliable` (support >= MIN_RELIABLE_SUPPORT)
                  so an unmeasurable class reads as unmeasurable.
- Confidence   -- accuracy by self-reported confidence band, per field.
                  Accuracy should climb with confidence; if it does not, the
                  number is decoration and cannot be a downstream filter.
- Cases        -- one row per case x field: gold, model answer, correct,
                  confidence, plus (from the case file) the model draft the
                  gold started from and whether the reviewer changed it.
- Adjudicate   -- the disagreements only, with empty `verdict` / `notes`
                  columns to fill in by hand. The share of rows where the
                  gold turns out to be wrong is the number that says whether
                  the realistic ceiling is 85% or 95%.

No model calls, no Langfuse; reads only the report JSON and the case files.
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from scripts.profile.business_profile_eval import case_files, load_case
from scripts.profile.business_profile_metrics import MIN_RELIABLE_SUPPORT, SCORED_FIELDS

DEFAULT_CASES_DIR = Path("evals/business_profiles/cases")

HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
WRONG_FILL = PatternFill("solid", fgColor="F8D7DA")
RIGHT_FILL = PatternFill("solid", fgColor="D4EDDA")

FIELD_METRICS: tuple[tuple[str, str], ...] = (
    ("accuracy", "accuracy"),
    ("majority_baseline", "majority baseline"),
    ("lift_over_baseline", "lift over baseline"),
    ("coverage", "coverage (committed, not unclear)"),
    ("accuracy_when_committed_on_answerable", "accuracy when committed, answerable cases"),
    ("macro_f1", "macro F1 (present classes)"),
    ("macro_f1_reliable_classes_only", f"macro F1 (classes with >= {MIN_RELIABLE_SUPPORT} gold)"),
    ("scored", "n scored"),
    ("abstained", "n unclear"),
    ("rejected", "n rejected"),
)


def run_label(report: dict[str, Any]) -> str:
    """How a run is named in column headers: model, prompt version, date."""
    when = report.get("generated_at") or ""
    try:
        day = datetime.fromisoformat(when).strftime("%Y-%m-%d")
    except ValueError:
        day = when[:10]
    label = f"{report.get('model', '?')} @ {report.get('prompt_version', '?')} ({day})"
    # A rescore of saved responses carries the *current* prompt version
    # (that is the rule set it was scored under); the responses themselves
    # came from an older run, and the header must say so or two columns
    # of one comparison read as the same run.
    rescored_from = report.get("rescored_from")
    if rescored_from:
        source = Path(str(rescored_from)).name
        label = f"{label} [rescore of {source}]"
    return label


def load_report(path: Path) -> dict[str, Any]:
    report = json.loads(path.read_text(encoding="utf-8"))
    report["_path"] = str(path)
    return report


def case_index(cases_dir: Path) -> dict[str, dict[str, Any]]:
    """company_number -> case, for names, drafts, and review decisions."""
    if not cases_dir.exists():
        return {}
    return {case["company_number"]: case for case in (load_case(p) for p in case_files(cases_dir))}


# --- tab builders -------------------------------------------------------------

def _header(ws: Worksheet, row: int, values: list[Any]) -> None:
    for col, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _autosize(ws: Worksheet, max_width: int = 60) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max_width, max(10, width + 2))


def summary_rows(reports: list[dict[str, Any]]) -> list[list[Any]]:
    """The Summary tab as rows: a label column then one column per run."""
    rows: list[list[Any]] = [["metric", *(run_label(r) for r in reports)]]

    def add(label: str, pick: Any) -> None:
        rows.append([label, *(pick(r) for r in reports)])

    rows.append(["run"])
    add("report file", lambda r: r.get("_path"))
    add("config", lambda r: r.get("config"))
    add("cases", lambda r: r.get("cases"))
    add("responses rejected outright (no JSON)", lambda r: r.get("responses_rejected_outright", r.get("quote_or_validation_rejections")))
    add("responses with a dropped field", lambda r: r.get("responses_with_dropped_fields"))
    add("fields dropped (of cases x 6)", lambda r: r.get("fields_rejected"))
    add("field pass rate", lambda r: r.get("field_pass_rate"))
    add("responses touched by any rejection", lambda r: r.get("quote_or_validation_rejections"))
    add("unclear rate (all fields)", lambda r: r.get("unclear_rate"))
    add("mean field accuracy", lambda r: (r.get("metrics") or {}).get("mean_field_accuracy"))
    add("elapsed seconds", lambda r: r.get("elapsed_seconds"))

    rows.append(["search-addressable (can paid search reach this company?)"])
    for key in ("precision", "recall", "f1", "tp", "fp", "fn", "gold_positives", "considered",
                "missed_by_abstention", "floor_rescued_gold", "floor_rescued_predicted"):
        add(key.replace("_", " "), lambda r, k=key: ((r.get("metrics") or {}).get("search_addressable") or {}).get(k))

    for field in SCORED_FIELDS:
        rows.append([field])
        for key, label in FIELD_METRICS:
            add(f"  {label}", lambda r, f=field, k=key: (((r.get("metrics") or {}).get("fields") or {}).get(f) or {}).get(k))
        add("  classes too thin to measure", lambda r, f=field: ", ".join(
            (((r.get("metrics") or {}).get("fields") or {}).get(f) or {}).get("classes_below_min_support") or []
        ) or None)
    return rows


def per_class_rows(reports: list[dict[str, Any]]) -> list[list[Any]]:
    rows: list[list[Any]] = [["run", "field", "value", "gold support", "reliable", "precision", "recall", "f1", "tp", "fp", "fn"]]
    for report in reports:
        label = run_label(report)
        for field, metrics in ((report.get("metrics") or {}).get("fields") or {}).items():
            for value, entry in (metrics.get("per_class") or {}).items():
                rows.append([label, field, value, entry.get("support"), entry.get("reliable"),
                             entry.get("precision"), entry.get("recall"), entry.get("f1"),
                             entry.get("tp"), entry.get("fp"), entry.get("fn")])
    return rows


def confidence_rows(reports: list[dict[str, Any]]) -> list[list[Any]]:
    rows: list[list[Any]] = [["run", "field", "confidence band", "n committed (answerable)", "accuracy", "answers with no confidence"]]
    for report in reports:
        label = run_label(report)
        for field, metrics in ((report.get("metrics") or {}).get("fields") or {}).items():
            bands = metrics.get("confidence_bands")
            if not bands:
                continue
            for band in bands.get("bands") or []:
                lo, hi = band["range"]
                rows.append([label, field, f"{lo:.2f} - {hi:.2f}", band.get("support"), band.get("accuracy"),
                             bands.get("missing_confidence")])
    return rows


def _gold_provenance(case: dict[str, Any] | None, field: str) -> tuple[Any, str]:
    """(model draft value, how the gold got there) from the case file."""
    if not case:
        return None, "case file not found"
    draft = ((case.get("draft") or {}).get("expected") or {}).get(field)
    draft_value = draft.get("value") if isinstance(draft, dict) else None
    review = case.get("review") or {}
    if review.get("status") != "verified":
        return draft_value, f"not verified ({review.get('status')})"
    if field in (review.get("changed_fields") or []):
        return draft_value, "reviewer changed the draft"
    if draft is not None:
        return draft_value, "reviewer kept the draft"
    return draft_value, "reviewed (no draft on file)"


CASE_HEADER = ["run", "company number", "company name", "sic label", "field", "gold", "model answer",
               "correct", "model confidence", "model draft the gold started from", "gold provenance"]


def case_rows(reports: list[dict[str, Any]], cases: dict[str, dict[str, Any]]) -> list[list[Any]]:
    rows: list[list[Any]] = [CASE_HEADER]
    for report in reports:
        label = run_label(report)
        for result in report.get("results") or []:
            number = result.get("company_number")
            case = cases.get(number)
            for field, outcome in (result.get("fields") or {}).items():
                expected = outcome.get("expected")
                if expected is None:
                    continue  # no gold label: unscoreable, not wrong
                draft_value, provenance = _gold_provenance(case, field)
                rows.append([label, number, (case or {}).get("company_name"), (case or {}).get("sic_label"),
                             field, expected, outcome.get("actual"), bool(outcome.get("correct")),
                             outcome.get("confidence"), draft_value, provenance])
    return rows


ADJUDICATE_HEADER = [*CASE_HEADER, "verdict (gold right / model right / both defensible / neither)", "notes"]


def adjudicate_rows(rows: list[list[Any]]) -> list[list[Any]]:
    """The Cases rows the model got wrong, plus two blank columns to fill in."""
    correct_col = CASE_HEADER.index("correct")
    out: list[list[Any]] = [ADJUDICATE_HEADER]
    for row in rows[1:]:
        if not row[correct_col]:
            out.append([*row, None, None])
    return out


# --- workbook -----------------------------------------------------------------

ALL_TABS: tuple[str, ...] = ("Summary", "Per-class", "Confidence", "Cases", "Adjudicate")


def build_workbook(
    reports: list[dict[str, Any]],
    cases: dict[str, dict[str, Any]],
    *,
    compact: bool = False,
    tabs: tuple[str, ...] | None = None,
) -> Workbook:
    """``compact`` leaves out the Cases tab -- every row the model got wrong
    is still on Adjudicate with the same columns, and the full per-case
    listing is what makes the file too large to push through a size-limited
    upload (the Drive connector takes the file inline, base64-encoded).
    ``tabs`` picks an explicit subset of ALL_TABS instead."""
    wanted = set(tabs) if tabs else set(ALL_TABS) - ({"Cases"} if compact else set())
    unknown = wanted - set(ALL_TABS)
    if unknown:
        raise ValueError(f"unknown tab(s): {', '.join(sorted(unknown))}; choose from {', '.join(ALL_TABS)}")
    wb = Workbook()
    wb.remove(wb.active)

    def add_sheet(title: str, rows: list[list[Any]], *, highlight_correct_col: int | None = None) -> Worksheet:
        ws = wb.create_sheet(title)
        _header(ws, 1, rows[0])
        for r, row in enumerate(rows[1:], 2):
            for c, value in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=value)
                if isinstance(value, float):
                    cell.number_format = "0.000"
            if highlight_correct_col is not None and len(row) > highlight_correct_col:
                ws.cell(row=r, column=highlight_correct_col + 1).fill = (
                    RIGHT_FILL if row[highlight_correct_col] else WRONG_FILL
                )
        ws.freeze_panes = "B2" if title == "Summary" else "A2"
        _autosize(ws)
        return ws

    if "Summary" in wanted:
        summary = add_sheet("Summary", summary_rows(reports))
        for row in summary.iter_rows(min_row=2):
            if row[0].value and not str(row[0].value).startswith("  ") and all(c.value is None for c in row[1:]):
                row[0].font = Font(bold=True)  # section heading rows
    if "Per-class" in wanted:
        add_sheet("Per-class", per_class_rows(reports))
    if "Confidence" in wanted:
        add_sheet("Confidence", confidence_rows(reports))
    rows = case_rows(reports, cases)
    if "Cases" in wanted:
        add_sheet("Cases", rows, highlight_correct_col=CASE_HEADER.index("correct"))
    if "Adjudicate" in wanted:
        add_sheet("Adjudicate", adjudicate_rows(rows), highlight_correct_col=CASE_HEADER.index("correct"))
    return wb


def write_csv(rows: list[list[Any]], path: Path) -> None:
    """One tab as CSV -- plain text survives a size-limited or hand-copied
    upload where a binary workbook would not."""
    import csv

    with path.open("w", encoding="utf-8", newline="") as fh:
        csv.writer(fh).writerows(rows)


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("reports", nargs="+", type=Path, help="Report JSON(s) from `business_profile_eval run`.")
    parser.add_argument("--out", type=Path, default=None,
                        help="Workbook path. Default: logs/business-profile-eval/eval-sheet-<timestamp>.xlsx")
    parser.add_argument("--cases-dir", type=Path, default=DEFAULT_CASES_DIR)
    parser.add_argument("--compact", action="store_true",
                        help="Omit the Cases tab (Adjudicate keeps every miss). For publishing; keep the full file too.")
    parser.add_argument("--tabs", default=None, help=f"Comma-separated subset of tabs: {', '.join(ALL_TABS)}.")
    parser.add_argument("--adjudicate-csv", type=Path, default=None,
                        help="Also write the Adjudicate rows as a CSV at this path.")
    args = parser.parse_args(argv)

    reports = [load_report(path) for path in args.reports]
    out = args.out or Path("logs/business-profile-eval") / f"eval-sheet-{datetime.now().strftime('%Y%m%dT%H%M%S')}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    cases = case_index(args.cases_dir)
    tabs = tuple(t.strip() for t in args.tabs.split(",")) if args.tabs else None
    build_workbook(reports, cases, compact=args.compact, tabs=tabs).save(out)
    print(f"Workbook written to {out}")
    if args.adjudicate_csv:
        write_csv(adjudicate_rows(case_rows(reports, cases)), args.adjudicate_csv)
        print(f"Adjudicate CSV written to {args.adjudicate_csv}")
    for report in reports:
        print(f"  {run_label(report)}: {report.get('cases')} cases")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from scripts.profile import business_profile_report_sheet as S
from scripts.profile.business_profile_metrics import compute_metrics, score_case


def _case(number: str, demand: str, changed: bool) -> dict:
    draft_demand = "b2b_relationship" if changed else demand
    return {
        "company_number": number,
        "company_name": f"COMPANY {number}",
        "sic_label": "Widgets",
        "sections": {"filed_report": "text"},
        "draft": {"expected": {"demand_model": {"value": draft_demand}, "customer_type": {"value": "b2c"}}},
        "expected": {"demand_model": {"value": demand}, "customer_type": {"value": "b2c"}},
        "review": {"status": "verified", "changed_fields": ["demand_model"] if changed else []},
    }


def _report(model: str, cases: list[dict], answers: dict[str, str]) -> dict:
    results = []
    for case in cases:
        extracted = {
            "demand_model": {"value": answers[case["company_number"]], "confidence": 0.8},
            "customer_type": {"value": "b2c", "confidence": 0.95},
        }
        results.append(score_case(case, extracted, scored_fields=("demand_model", "customer_type")))
    metrics = compute_metrics(results, scored_fields=("demand_model", "customer_type"))
    return {
        "generated_at": "2026-09-12T10:00:00+00:00",
        "config": "x.yaml",
        "model": model,
        "prompt_version": "business-profile-v6",
        "cases": len(cases),
        "quote_or_validation_rejections": 0,
        "quote_verification_pass_rate": 1.0,
        "unclear_rate": 0.0,
        "elapsed_seconds": 1.0,
        "metrics": metrics,
        "results": results,
    }


def test_workbook_has_every_tab_and_compares_runs_side_by_side(tmp_path: Path) -> None:
    cases = [_case("1", "consumer_search", changed=True), _case("2", "local_service", changed=False)]
    a = _report("model-a", cases, {"1": "consumer_search", "2": "b2b_relationship"})
    b = _report("model-b", cases, {"1": "unclear", "2": "local_service"})
    wb = S.build_workbook([a, b], {c["company_number"]: c for c in cases})
    out = tmp_path / "sheet.xlsx"
    wb.save(out)

    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Per-class", "Confidence", "Cases", "Adjudicate"]

    summary = wb["Summary"]
    header = [c.value for c in summary[1]]
    assert header == ["metric", S.run_label(a), S.run_label(b)]
    rows = [[c.value for c in row] for row in summary.iter_rows(min_row=2)]
    by_label = {row[0]: row[1:] for row in rows}
    assert by_label["cases"] == [2, 2]
    # per-field metric labels repeat under each field heading: read demand_model's block
    start = next(i for i, row in enumerate(rows) if row[0] == "demand_model")
    accuracy = next(row for row in rows[start:] if row[0] == "  accuracy")
    assert accuracy[1] == a["metrics"]["fields"]["demand_model"]["accuracy"]
    assert by_label["precision"] == [
        a["metrics"]["search_addressable"]["precision"],
        b["metrics"]["search_addressable"]["precision"],
    ]


def test_cases_tab_carries_gold_provenance_and_adjudicate_lists_only_misses(tmp_path: Path) -> None:
    cases = [_case("1", "consumer_search", changed=True), _case("2", "local_service", changed=False)]
    report = _report("model-a", cases, {"1": "b2b_relationship", "2": "local_service"})
    wb = S.build_workbook([report], {c["company_number"]: c for c in cases})

    rows = [[c.value for c in row] for row in wb["Cases"].iter_rows()]
    assert rows[0] == S.CASE_HEADER
    demand_1 = next(r for r in rows[1:] if r[1] == "1" and r[4] == "demand_model")
    assert demand_1[5:8] == ["consumer_search", "b2b_relationship", False]
    assert demand_1[9:11] == ["b2b_relationship", "reviewer changed the draft"]
    demand_2 = next(r for r in rows[1:] if r[1] == "2" and r[4] == "demand_model")
    assert demand_2[7] is True
    assert demand_2[10] == "reviewer kept the draft"

    adj = [[c.value for c in row] for row in wb["Adjudicate"].iter_rows()]
    assert adj[0][-2].startswith("verdict")
    assert [r[1:5] for r in adj[1:]] == [["1", "COMPANY 1", "Widgets", "demand_model"]]
    assert adj[1][-2:] == [None, None]  # left for the human


def test_main_writes_the_workbook_where_asked(tmp_path: Path, monkeypatch) -> None:
    cases = [_case("1", "consumer_search", changed=False)]
    cases_dir = tmp_path / "cases"
    cases_dir.mkdir()
    (cases_dir / "1.json").write_text(json.dumps(cases[0]), encoding="utf-8")
    report_path = tmp_path / "report.json"
    report_path.write_text(json.dumps(_report("m", cases, {"1": "consumer_search"})), encoding="utf-8")
    out = tmp_path / "out.xlsx"

    assert S.main([str(report_path), "--out", str(out), "--cases-dir", str(cases_dir)]) == 0
    assert load_workbook(out)["Cases"].max_row == 3  # header + 2 fields
